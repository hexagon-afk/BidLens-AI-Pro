"""
Multi-Format Document Processing & OCR Engine - Layer 3, Branch A
Extracts text, entities, and clause evidence from PDF, DOCX, XLSX, and scanned images.
"""
import pymupdf
import docx
import openpyxl
import pandas as pd
import re
import os
import cv2
import numpy as np

_ocr_engine = None

def get_ocr_engine():
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
            _ocr_engine = RapidOCR()
        except Exception as e:
            print(f"RapidOCR initialization notice: {e}")
            _ocr_engine = None
    return _ocr_engine


def extract_document_text(file_path: str) -> tuple[str, int, str]:
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    full_text = ""
    page_count = 1

    if ext == ".pdf":
        try:
            doc = pymupdf.open(file_path)
            page_count = len(doc)
            digital_pages = []
            pages_needing_ocr = []

            for i, page in enumerate(doc):
                t = page.get_text()
                if t and len(t.strip()) > 30:
                    digital_pages.append(f"\n--- Page {i + 1} ---\n" + t)
                else:
                    pages_needing_ocr.append(i)

            if len(digital_pages) > 0 and len("".join(digital_pages).strip()) >= 50:
                full_text = "\n".join(digital_pages)
            else:
                ocr = get_ocr_engine()
                ocr_text_parts = []
                if ocr:
                    for i in range(page_count):
                        page = doc[i]
                        pix = page.get_pixmap(dpi=150)
                        img_bytes = pix.tobytes("png")
                        res, _ = ocr(img_bytes)
                        if res:
                            lines = [r[1] for r in res]
                            ocr_text_parts.append(f"\n--- Page {i + 1} (OCR) ---\n" + "\n".join(lines))
                full_text = "\n".join(ocr_text_parts) if ocr_text_parts else "--- Scanned PDF ---"
            doc.close()
        except Exception as pdf_err:
            full_text = f"Error: {pdf_err}"

    elif ext in [".jpg", ".jpeg", ".png"]:
        ocr = get_ocr_engine()
        page_count = 1
        if ocr:
            res, _ = ocr(file_path)
            if res:
                full_text = "\n".join([r[1] for r in res])

    elif ext in [".docx", ".doc"]:
        doc = docx.Document(file_path)
        full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        page_count = max(1, len(full_text) // 1500)

    elif ext in [".xlsx", ".xls"]:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet_texts = []
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            for row in sheet.iter_rows(values_only=True):
                row_vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if row_vals:
                    sheet_texts.append(" | ".join(row_vals))
        full_text = "\n".join(sheet_texts)
        page_count = len(wb.sheetnames)

    return full_text, page_count, ext.replace(".", "").upper()


def extract_document_data(file_path: str) -> dict:
    full_text, page_count, file_type = extract_document_text(file_path)

    gstin_matches = re.findall(r"\b\d{2}[A-Z]{5}\d{4}[A-Z]{1}[A-Z\d]{1}[Z]{1}[A-Z\d]{1}\b", full_text)
    pan_matches = re.findall(r"\b[A-Z]{5}\d{4}[A-Z]{1}\b", full_text)
    udyam_matches = re.findall(r"\bUDYAM-[A-Z]{2}-\d{2}-\d{7}\b", full_text)

    turnover_cr = None
    to_match = re.search(r"(?:Annual\s+Turnover|Turnover)[^\d]*([\d.]+)\s*(crore|cr|lakh|lakhs)", full_text, re.IGNORECASE)
    if to_match:
        val = float(to_match.group(1))
        turnover_cr = val if "cr" in to_match.group(2).lower() else (val / 100.0)

    emd_status = "MISSING"
    if any(k in full_text.lower() for k in ["bank guarantee", "bg no", "fdr", "1,00,000"]):
        emd_status = "SUBMITTED"
    elif udyam_matches or "msme" in full_text.lower():
        emd_status = "MSME_EXEMPT"

    warranty_terms = "Standard"
    if "5-year" in full_text.lower() or "5 year" in full_text.lower():
        warranty_terms = "5-Year Comprehensive 24x7 Onsite Warranty"
    elif "3-year" in full_text.lower():
        warranty_terms = "3-Year Comprehensive Warranty"

    return {
        "filename": os.path.basename(file_path),
        "file_type": file_type,
        "vendor_name": "Vendor Proposal",
        "page_count": page_count,
        "gstin": gstin_matches[0] if gstin_matches else None,
        "all_gstins": gstin_matches,
        "gstin_expired": "EXPIRED" in full_text.upper() or "CANCELLED" in full_text.upper(),
        "pan": pan_matches[0] if pan_matches else None,
        "all_pans": list(set(pan_matches)),
        "udyam": udyam_matches[0] if udyam_matches else None,
        "is_msme": len(udyam_matches) > 0 or "msme" in full_text.lower(),
        "turnover_cr": turnover_cr,
        "emd_status": emd_status,
        "warranty": warranty_terms,
        "raw_text": full_text[:1200]
    }
